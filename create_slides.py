import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Create presentation
prs = Presentation()

# Set slide width and height to 1920x1080 (20x11.25 inches)
prs.slide_width = Inches(20)
prs.slide_height = Inches(11.25)

# Add a blank slide
blank_slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_slide_layout)

# Define Google Photos colors
BLUE = RGBColor(0x42, 0x85, 0xF4)
RED = RGBColor(0xEA, 0x43, 0x35)
YELLOW = RGBColor(0xFB, 0xBC, 0x05)
GREEN = RGBColor(0x34, 0xA8, 0x53)
DARK_GRAY = RGBColor(0x20, 0x21, 0x24)
LIGHT_GRAY = RGBColor(0xF1, 0xF3, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Set background color to very light gray / off-white
background = slide.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0xF8, 0xF9, 0xFA)

def add_text_box(slide, text, left, top, width, height, font_size=22, bold=False, color=DARK_GRAY, align=PP_ALIGN.LEFT, bg_color=None):
    if bg_color:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = bg_color
    else:
        shape = slide.shapes.add_textbox(left, top, width, height)
    
    text_frame = shape.text_frame
    text_frame.word_wrap = True
    p = text_frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(font_size)
    p.font.name = 'Calibri'
    p.font.bold = bold
    p.font.color.rgb = color
    return shape

# --- Header ---
title_text = "Google Photos | Core Experience | Improving Retrieval of Vaguely Remembered Photos"
add_text_box(slide, title_text, Inches(1), Inches(0.5), Inches(18), Inches(1), font_size=32, bold=True, color=DARK_GRAY, align=PP_ALIGN.LEFT)

subtitle_text = "Business Metric Decomposition"
add_text_box(slide, subtitle_text, Inches(11), Inches(0.5), Inches(8), Inches(1), font_size=32, bold=True, color=DARK_GRAY, align=PP_ALIGN.RIGHT)


# --- LEFT SIDE: Metrics and Strategy ---
# 4 Boxes for Metrics
box_w = Inches(4.2)
box_h = Inches(2)
box_y1 = Inches(2)
box_y2 = Inches(4.3)
box_x1 = Inches(1)
box_x2 = Inches(5.5)

# Box 1 (Blue)
shape1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, box_x1, box_y1, box_w, box_h)
shape1.fill.solid()
shape1.fill.fore_color.rgb = BLUE
shape1.line.color.rgb = BLUE
tf1 = shape1.text_frame
tf1.word_wrap = True
p1 = tf1.paragraphs[0]
p1.text = "2B+"
p1.font.size = Pt(40)
p1.font.bold = True
p1.font.name = 'Calibri'
p1.font.color.rgb = WHITE
p1.alignment = PP_ALIGN.CENTER
p1_sub = tf1.add_paragraph()
p1_sub.text = "Monthly Active Users"
p1_sub.font.size = Pt(24)
p1_sub.font.name = 'Calibri'
p1_sub.font.color.rgb = WHITE
p1_sub.alignment = PP_ALIGN.CENTER

# Box 2 (Red)
shape2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, box_x2, box_y1, box_w, box_h)
shape2.fill.solid()
shape2.fill.fore_color.rgb = RED
shape2.line.color.rgb = RED
tf2 = shape2.text_frame
tf2.word_wrap = True
p2 = tf2.paragraphs[0]
p2.text = "350M+"
p2.font.size = Pt(40)
p2.font.bold = True
p2.font.name = 'Calibri'
p2.font.color.rgb = WHITE
p2.alignment = PP_ALIGN.CENTER
p2_sub = tf2.add_paragraph()
p2_sub.text = "Google Subscriptions (Q1 2026)"
p2_sub.font.size = Pt(22)
p2_sub.font.name = 'Calibri'
p2_sub.font.color.rgb = WHITE
p2_sub.alignment = PP_ALIGN.CENTER

# Box 3 (Yellow)
shape3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, box_x1, box_y2, box_w, box_h)
shape3.fill.solid()
shape3.fill.fore_color.rgb = YELLOW
shape3.line.color.rgb = YELLOW
tf3 = shape3.text_frame
tf3.word_wrap = True
p3 = tf3.paragraphs[0]
p3.text = "370M+"
p3.font.size = Pt(40)
p3.font.bold = True
p3.font.name = 'Calibri'
p3.font.color.rgb = WHITE
p3.alignment = PP_ALIGN.CENTER
p3_sub = tf3.add_paragraph()
p3_sub.text = "Photo Searches Monthly"
p3_sub.font.size = Pt(24)
p3_sub.font.name = 'Calibri'
p3_sub.font.color.rgb = WHITE
p3_sub.alignment = PP_ALIGN.CENTER

# Box 4 (Green)
shape4 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, box_x2, box_y2, box_w, box_h)
shape4.fill.solid()
shape4.fill.fore_color.rgb = GREEN
shape4.line.color.rgb = GREEN
tf4 = shape4.text_frame
tf4.word_wrap = True
p4 = tf4.paragraphs[0]
p4.text = "9T+"
p4.font.size = Pt(40)
p4.font.bold = True
p4.font.name = 'Calibri'
p4.font.color.rgb = WHITE
p4.alignment = PP_ALIGN.CENTER
p4_sub = tf4.add_paragraph()
p4_sub.text = "Total Photos Stored"
p4_sub.font.size = Pt(24)
p4_sub.font.name = 'Calibri'
p4_sub.font.color.rgb = WHITE
p4_sub.alignment = PP_ALIGN.CENTER

# Strategic Objective Box
strat_y = Inches(6.6)
strat_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), strat_y, Inches(8.7), Inches(2))
strat_shape.fill.solid()
strat_shape.fill.fore_color.rgb = RGBColor(0xE8, 0xF0, 0xFE) # Light blue
strat_shape.line.color.rgb = BLUE
strat_tf = strat_shape.text_frame
strat_tf.word_wrap = True
strat_p1 = strat_tf.paragraphs[0]
strat_p1.text = "Google Photos Strategic Objective:"
strat_p1.font.size = Pt(26)
strat_p1.font.bold = True
strat_p1.font.name = 'Calibri'
strat_p1.font.color.rgb = DARK_GRAY
strat_p1.alignment = PP_ALIGN.CENTER
strat_p2 = strat_tf.add_paragraph()
strat_p2.text = "Increase the percentage of users who successfully retrieve a photo they remember but cannot precisely describe when they start searching."
strat_p2.font.size = Pt(22)
strat_p2.font.name = 'Calibri'
strat_p2.font.color.rgb = DARK_GRAY
strat_p2.alignment = PP_ALIGN.CENTER

# Why it matters Box
why_y = Inches(8.8)
why_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), why_y, Inches(8.7), Inches(1.8))
why_shape.fill.solid()
why_shape.fill.fore_color.rgb = RGBColor(0xFE, 0xEF, 0xC3) # Light yellow
why_shape.line.color.rgb = YELLOW
why_tf = why_shape.text_frame
why_tf.word_wrap = True
why_p1 = why_tf.paragraphs[0]
why_p1.text = "Why this strategic objective matters?"
why_p1.font.size = Pt(24)
why_p1.font.bold = True
why_p1.font.name = 'Calibri'
why_p1.font.color.rgb = DARK_GRAY
why_p1.alignment = PP_ALIGN.CENTER
why_p2 = why_tf.add_paragraph()
why_p2.text = "Improving retrieval of vaguely remembered photos increases user trust. When users know they can find any memory effortlessly, they rely on Google Photos as their primary vault, organically driving Google One subscription upgrades due to storage limits."
why_p2.font.size = Pt(22)
why_p2.font.name = 'Calibri'
why_p2.font.color.rgb = DARK_GRAY
why_p2.alignment = PP_ALIGN.CENTER

# Sources
add_text_box(slide, "Sources: Business Insider (2026), Business of Apps (2026), PetaPixel (2026)", Inches(1), Inches(10.7), Inches(10), Inches(0.5), font_size=18, color=RGBColor(0x5F, 0x63, 0x68))

# --- RIGHT SIDE: Flowchart ---
from pptx.enum.shapes import MSO_CONNECTOR

def add_opportunity_box(slide, left, top, width, height, border_color, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = border_color
    shape.line.width = Pt(2)
    
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(16)
    p.font.name = 'Calibri'
    p.font.color.rgb = DARK_GRAY
    p.alignment = PP_ALIGN.CENTER
    return shape

def add_connector(slide, start_x, start_y, end_x, end_y):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, start_x, start_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0x9A, 0x9C, 0xA0)
    connector.line.width = Pt(2)

# Top Node
top_w = Inches(4)
top_h = Inches(1.5)
top_x = Inches(12.75)
top_y = Inches(1.5)
add_opportunity_box(slide, top_x, top_y, top_w, top_h, BLUE, "1. Metadata Dependency\n(Score: 68.9)")

# Middle Nodes
mid_w = Inches(3.5)
mid_h = Inches(1.5)
mid_y = Inches(4.5)
mid_x1 = Inches(10.75)
mid_x2 = Inches(15.25)
add_opportunity_box(slide, mid_x1, mid_y, mid_w, mid_h, RED, "2. Search Strategy Failure\n(Score: 62.8)")
add_opportunity_box(slide, mid_x2, mid_y, mid_w, mid_h, YELLOW, "3. Result Ranking Failure\n(Score: 61.9)")

# Connect Top to Middle
add_connector(slide, top_x + top_w/2, top_y + top_h, mid_x1 + mid_w/2, mid_y)
add_connector(slide, top_x + top_w/2, top_y + top_h, mid_x2 + mid_w/2, mid_y)

# Bottom Nodes
bot_w = Inches(3.5)
bot_h = Inches(1.5)
bot_y = Inches(7.5)
bot_x1 = Inches(10.75)
bot_x2 = Inches(15.25)
add_opportunity_box(slide, bot_x1, bot_y, bot_w, bot_h, GREEN, "4. Context Loss\n(Score: 54.8)")
add_opportunity_box(slide, bot_x2, bot_y, bot_w, bot_h, GREEN, "5. Person Ambiguity\n(Score: 54.5)")

# Connect Middle to Bottom
add_connector(slide, mid_x1 + mid_w/2, mid_y + mid_h, bot_x1 + bot_w/2, bot_y)
add_connector(slide, mid_x2 + mid_w/2, mid_y + mid_h, bot_x2 + bot_w/2, bot_y)

prs.save('/Users/raunaqkaicker/Documents/Google photos AI discovery engine/Google_Photos_Metric_Decomposition.pptx')
