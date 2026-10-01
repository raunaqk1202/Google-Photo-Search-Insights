import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_slide():
    prs = Presentation()
    # 1920x1080 pixels at 96 DPI is 20 x 11.25 inches
    prs.slide_width = Inches(20)
    prs.slide_height = Inches(11.25)

    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    # Define Google Colors
    BLUE = RGBColor(66, 133, 244)
    RED = RGBColor(234, 67, 53)
    YELLOW = RGBColor(251, 188, 4)
    GREEN = RGBColor(52, 168, 83)
    DARK_GRAY = RGBColor(64, 64, 64)
    LIGHT_GRAY = RGBColor(245, 245, 245)
    WHITE = RGBColor(255, 255, 255)

    def add_textbox(slide, text, left, top, width, height, font_size=22, bold=False, color=DARK_GRAY, alignment=PP_ALIGN.LEFT):
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = alignment
        p.font.size = Pt(font_size)
        p.font.name = 'Calibri'
        p.font.bold = bold
        p.font.color.rgb = color
        return txBox

    # Adding colorful squares in the top-left like the screenshot
    sq_size = Inches(0.8)
    s1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0), sq_size, sq_size)
    s1.fill.solid(); s1.fill.fore_color.rgb = BLUE; s1.line.fill.background()

    s2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), sq_size, sq_size, sq_size)
    s2.fill.solid(); s2.fill.fore_color.rgb = RED; s2.line.fill.background()

    s3 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, sq_size, sq_size * 1.8, sq_size, sq_size)
    s3.fill.solid(); s3.fill.fore_color.rgb = YELLOW; s3.line.fill.background()

    # Green square on right
    s4 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(18.5), Inches(1.5), sq_size, sq_size)
    s4.fill.solid(); s4.fill.fore_color.rgb = GREEN; s4.line.fill.background()

    # Add Title
    add_textbox(slide, "Google Photos AI Discovery Engine", Inches(2), Inches(0.3), Inches(16), Inches(0.8), font_size=44, bold=True, color=DARK_GRAY, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, "Engine Workflow, Scoring Algorithm, and Key Opportunities", Inches(2), Inches(1.0), Inches(16), Inches(0.5), font_size=26, bold=False, color=DARK_GRAY, alignment=PP_ALIGN.CENTER)

    # --- STATS SECTION ---
    def draw_stat_card(x, y, w, h, top_color, text, subtext):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shape.fill.solid(); shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = top_color
        shape.line.width = Pt(2)
        
        tf = shape.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = text
        p1.font.bold = True
        p1.font.size = Pt(28)
        p1.font.color.rgb = DARK_GRAY
        p1.alignment = PP_ALIGN.CENTER
        
        p2 = tf.add_paragraph()
        p2.text = subtext
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = DARK_GRAY
        p2.alignment = PP_ALIGN.CENTER

    # Row 1
    draw_stat_card(Inches(1), Inches(1.6), Inches(8.5), Inches(0.8), BLUE, "1,982", "Reviews Scraped")
    draw_stat_card(Inches(10.5), Inches(1.6), Inches(8.5), Inches(0.8), RGBColor(236, 72, 153), "507", "Classified Reviews") # Pink

    # Row 2
    draw_stat_card(Inches(1), Inches(2.5), Inches(3.2), Inches(0.7), GREEN, "114", "Play Store")
    draw_stat_card(Inches(4.6), Inches(2.5), Inches(3.2), Inches(0.7), RED, "800", "YouTube")
    draw_stat_card(Inches(8.2), Inches(2.5), Inches(3.2), Inches(0.7), BLUE, "32", "App Store")
    draw_stat_card(Inches(11.8), Inches(2.5), Inches(3.2), Inches(0.7), RGBColor(249, 115, 22), "826", "Reddit") # Orange
    draw_stat_card(Inches(15.4), Inches(2.5), Inches(3.2), Inches(0.7), YELLOW, "210", "Google Support")

    # 1. Flowchart (Left Column)
    flowchart_x = Inches(1)
    add_textbox(slide, "Engine Workflow", flowchart_x, Inches(3.4), Inches(5.5), Inches(0.5), font_size=28, bold=True, color=BLUE)

    steps = [
        ("1. Data Scraping", "Scrape user feedback (Play Store, Reddit, etc.)"),
        ("2. Data Structuring", "Clean and format unstructured text data"),
        ("3. AI Analysis", "LLM categorization of failure modes"),
        ("4. Database Storage", "PostgreSQL storage with vector embeddings"),
        ("5. RAG Pipeline", "Groq-powered conversational UX analytics")
    ]

    for i, (title, desc) in enumerate(steps):
        y = Inches(4.1 + i * 1.3)
        # add box shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, flowchart_x, y, Inches(5.5), Inches(0.9))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = BLUE
        shape.line.width = Pt(2)
        
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.bold = True
        p1.font.size = Pt(20)
        p1.font.color.rgb = DARK_GRAY
        p1.alignment = PP_ALIGN.CENTER
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(18)
        p2.font.color.rgb = DARK_GRAY
        p2.alignment = PP_ALIGN.CENTER

        if i < len(steps) - 1:
            # add arrow down
            arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, flowchart_x + Inches(2.55), y + Inches(0.95), Inches(0.4), Inches(0.3))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = BLUE
            arrow.line.fill.background()

    # 2. Scoring Algorithm (Middle Column)
    algo_x = Inches(7)
    add_textbox(slide, "Scoring Algorithm", algo_x, Inches(3.4), Inches(5.5), Inches(0.5), font_size=28, bold=True, color=RED)

    algo_text = (
        "Prioritizes product opportunities based on a 0-100 score:\n\n"
        "Component Weights:\n"
        "• 35% Reach Score (estimated frequency)\n"
        "• 30% User Pain Score\n"
        "• 20% Business Impact Score\n"
        "• 15% Evidence Strength\n\n"
        "Extraction Method:\n"
        "All component scores (1.0-5.0) are evaluated and assigned by the LLM based on user feedback.\n\n"
        "Final Score Calculation:\n"
        "Raw Score = Weighted sum of components\n"
        "Score = (Raw Score - 1.0) * 25"
    )
    algo_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, algo_x, Inches(4.1), Inches(5.5), Inches(6.1))
    algo_box.fill.solid()
    algo_box.fill.fore_color.rgb = WHITE
    algo_box.line.color.rgb = RED
    algo_box.line.width = Pt(2)

    tf_algo = algo_box.text_frame
    tf_algo.word_wrap = True
    tf_algo.margin_left = Inches(0.2)
    tf_algo.margin_top = Inches(0.2)
    tf_algo.margin_right = Inches(0.2)
    
    first = True
    for line in algo_text.split('\n'):
        if first:
            p = tf_algo.paragraphs[0]
            first = False
        else:
            p = tf_algo.add_paragraph()
            
        p.text = line
        p.font.size = Pt(19)
        p.font.name = 'Calibri'
        p.font.color.rgb = DARK_GRAY
        if "Component Weights:" in line or "Extraction Method" in line or "Final Score Calculation:" in line:
            p.font.bold = True
            p.font.color.rgb = RED

    # 3. Opportunities (Right Column)
    opp_x = Inches(13)
    add_textbox(slide, "Important Opportunities", opp_x, Inches(3.4), Inches(6), Inches(0.5), font_size=28, bold=True, color=GREEN)

    opportunities = [
        ("Result Ranking Failure", "Relevant photos may exist but are difficult to recognize among results. The sheer volume of matches overwhelms the user's ability to spot their target."),
        ("Person Ambiguity", "The user remembers who was involved but cannot identify the person in a way the system can use. Sometimes faces are obscured, or the system hasn't tagged them correctly."),
        ("Context Loss", "The user remembers an event or story rather than searchable attributes of the image. The emotional context doesn't map well to objective image tags."),
        ("Search Strategy Failure", "The user does not know which search mechanism or filter to use. They might default to a simple keyword search when a combination of filters would be more effective."),
        ("Metadata Dependency", "Successful retrieval depends on information the user no longer remembers. Metadata like exact dates or locations are easily forgotten by users.")
    ]

    colors = [BLUE, RED, YELLOW, GREEN, BLUE]

    for i, (title, desc) in enumerate(opportunities):
        y = Inches(4.1 + i * 1.25)
        
        # number box
        num_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, opp_x, y, Inches(0.8), Inches(1.1))
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = colors[i]
        num_box.line.fill.background()
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = f"{i+1:02d}"
        p_num.font.size = Pt(28)
        p_num.font.name = 'Calibri'
        p_num.font.bold = True
        p_num.font.color.rgb = WHITE
        p_num.alignment = PP_ALIGN.CENTER
        
        # desc box
        desc_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, opp_x + Inches(0.8), y, Inches(5.2), Inches(1.1))
        desc_box.fill.solid()
        desc_box.fill.fore_color.rgb = LIGHT_GRAY
        desc_box.line.fill.background()
        tf_desc = desc_box.text_frame
        tf_desc.word_wrap = True
        tf_desc.margin_left = Inches(0.15)
        tf_desc.margin_top = Inches(0.1)
        tf_desc.margin_right = Inches(0.1)
        
        p_title = tf_desc.paragraphs[0]
        p_title.text = title
        p_title.font.bold = True
        p_title.font.size = Pt(20)
        p_title.font.name = 'Calibri'
        p_title.font.color.rgb = colors[i]
        
        p_desc = tf_desc.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(17)
        p_desc.font.name = 'Calibri'
        p_desc.font.color.rgb = DARK_GRAY

    prs.save("Discovery_Engine_Insights.pptx")
    print("Slide generated successfully.")

if __name__ == "__main__":
    create_slide()
